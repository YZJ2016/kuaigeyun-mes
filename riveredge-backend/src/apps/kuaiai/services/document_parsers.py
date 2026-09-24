"""KU-AI S3 文档解析路由（KR-D12 / 设计 §9）。

``extract_text(source_type, data)``：按来源类型路由到对应轻量解析包，
全部内存 bytes 解析——不扫服务器磁盘、不写未隔离临时目录语料。

失败关闭约定（不落乱码/半截文本）：

- ``doc``/``ppt`` 旧二进制格式无轻量纯 Python 解析器 → 明确
  ValidationError 拒绝，提示转 docx/pptx 重传；
- 魔数/签名不符（pdf 缺 ``%PDF`` 头、docx/pptx/xlsx 非 zip ``PK`` 头）
  → ValidationError；
- 解析包抛错、抽出文本 strip 后为空 → ValidationError。
"""

from __future__ import annotations

import io

from infra.exceptions.exceptions import ValidationError

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"

# 旧二进制格式：无轻量纯 Python 解析器，明确拒绝而非写半截文本
_LEGACY_BINARY_TARGETS = {"doc": "docx", "ppt": "pptx"}
# 纯文本直解（utf-8）
_TEXT_TYPES = {"md", "markdown", "txt", "text"}
# OOXML zip 容器类型（共享 PK 魔数校验）
_ZIP_TYPES = {"docx", "pptx", "xlsx"}


def _require_magic(data: bytes, magic: bytes, ext: str) -> None:
    if not data.startswith(magic):
        raise ValidationError(f"文件内容与 {ext} 格式签名不符")


def _pdf_text(data: bytes) -> str:
    _require_magic(data, _PDF_MAGIC, "pdf")
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(c) for c in row if c is not None and str(c).strip()
                ]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    finally:
        workbook.close()


def _decode_text(data: bytes) -> str:
    """utf-8 优先，解码失败回落 gb18030（覆盖 GBK 系中文文本）。"""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("gb18030")


def _route(ext: str, data: bytes) -> str:
    if ext == "pdf":
        return _pdf_text(data)
    if ext in _ZIP_TYPES:
        _require_magic(data, _ZIP_MAGIC, ext)
        if ext == "docx":
            return _docx_text(data)
        if ext == "pptx":
            return _pptx_text(data)
        return _xlsx_text(data)
    if ext == "xls":
        # openpyxl 只读 xlsx；真实 xls（OLE2）会在 load_workbook 内明确失败
        return _xlsx_text(data)
    if ext in _TEXT_TYPES:
        return _decode_text(data)
    raise ValidationError(f"不支持的文档格式：{ext or '空'}")


async def extract_text(source_type: str, data: bytes) -> str:
    """按 ``source_type`` 路由解析；任何失败一律 ValidationError。"""
    ext = (source_type or "").strip().lower().lstrip(".")
    if ext in _LEGACY_BINARY_TARGETS:
        raise ValidationError(
            f"格式 {ext} 暂不支持，请转换为 {_LEGACY_BINARY_TARGETS[ext]} 后重新上传"
        )
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise ValidationError("文档内容为空")
    try:
        text = _route(ext, bytes(data))
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError(f"{ext or '未知'} 文档解析失败") from exc
    if not (text or "").strip():
        raise ValidationError("解析结果为空文本")
    return text
