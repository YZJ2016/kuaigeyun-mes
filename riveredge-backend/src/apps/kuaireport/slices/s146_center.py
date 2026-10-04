"""星报表 P2：报表中心列表、KRP-D6 十张系统报表种入、后端全量 Excel。

执行数据只调用 spec 145 的 ``execute_report(tenant_id, report_id, filters)``。
本文件不查询数据集、不写 SQL 正文。
"""

from __future__ import annotations

import json
import re
import uuid
from contextlib import asynccontextmanager
from datetime import datetime
from io import BytesIO
from typing import Any, AsyncIterator

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from pydantic import BaseModel, Field

from apps.kuaireport.constants import REPORT_DATA_SOURCE_UUID
from core.api.deps.access import require_permission_codes
from core.api.deps.deps import get_current_tenant
from infra.api.deps.deps import get_current_user
from infra.domain.tenant_context import TenantContextError, get_current_tenant_id
from infra.exceptions.exceptions import NotFoundError, ValidationError
from infra.models.user import User

CLASSIFY_COLUMN_DEFAULT = "未分类"
CATEGORY_SYSTEM = "system"
CATEGORY_CUSTOM = "custom"
STATUS_DRAFT = "DRAFT"
STATUS_PUBLISHED = "PUBLISHED"
ALLOWED_CATEGORIES = frozenset({CATEGORY_SYSTEM, CATEGORY_CUSTOM})
FORBIDDEN_REPORT_NAMES = frozenset({"分析报表", "跨业务报表"})

# (code, name, classify)。销售/采购保持原分类，其余按拍板写死。
SYSTEM_REPORTS: tuple[tuple[str, str, str], ...] = (
    ("inv_ledger", "库存台账", "库存"),
    ("wo_tracking", "工单跟踪", "生产"),
    ("qc_pass_rate", "检验合格率", "质量"),
    ("perf_stats", "效能统计", "生产"),
    ("node_timeliness", "节点时效", "综合"),
    ("process_efficiency", "处理效率", "综合"),
    ("sales_trace", "销售全链路追踪", "销售"),
    ("purchase_trace", "采购全链路追踪", "采购"),
    ("material_trace", "物料全生命周期追踪", "物料"),
    ("biz_overview", "业务综合看板", "综合"),
)

# 全量 Excel 按 limit/offset 向 execute_report 翻页，直到返回行数不足一页。
FULL_EXPORT_PAGE_SIZE = 200

# Content-Disposition 文件名只允许 ASCII 安全字符；其余折叠成下划线。
_SAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._-]+")
# 单元格文本以这些字符开头时会被 Excel 当公式执行，导出时前置 ' 防注入。
_FORMULA_PREFIXES = ("=", "+", "-", "@")

_EXECUTE_MODULES = (
    "apps.kuaireport.services.execute",
    "apps.kuaireport.services.execute_service",
    "apps.kuaireport.services.execute_report",
)

_SECRET_KEY = re.compile(r"(?i)password|secret|token|cipher|jdbc|dsn|connection")
_SECRET_VALUE = re.compile(r"(?i)jdbc:|password\s*=|secret\s*=|token\s*=")
_PATH_VALUE = re.compile(r"^(?:[A-Za-z]:\\|\\\\|/)")

router = APIRouter(prefix="/reports", tags=["kuaireport-center"])


class FullExcelRequest(BaseModel):
    filters: dict[str, Any] = Field(default_factory=dict)


def require_tenant_id() -> int:
    tenant_id = get_current_tenant_id()
    if tenant_id is None:
        raise TenantContextError("无租户上下文，不读取报表表")
    return int(tenant_id)


def system_report_config(data_source_uuid: str | None) -> dict[str, Any]:
    """绑定只写 extra.data_source_uuid。不写 SQL，不写列公式。"""
    return {
        "fields": [],
        "filters": [],
        "extra": {
            REPORT_DATA_SOURCE_UUID: data_source_uuid,
            "uni_report": {
                "templateId": "kuaireportTable",
                "showIndexColumn": True,
                "showSummaryRow": True,
            },
        },
    }


def resolve_classify(raw: str | None) -> str:
    if raw is None or raw == "":
        return CLASSIFY_COLUMN_DEFAULT
    return raw


def resolve_execute_report():
    """定位 145 的 execute_report。未落地时抛错，不在本切片执行 SQL。"""
    import importlib

    for name in _EXECUTE_MODULES:
        try:
            module = importlib.import_module(name)
        except ImportError:
            continue
        fn = getattr(module, "execute_report", None)
        if callable(fn):
            return fn
    raise RuntimeError(
        "execute_report 未落地（spec 145）。全量 Excel 不在本切片另写查询。"
    )


_REPORT_LIST_COLUMNS = (
    "id, uuid, code, name, description, category, classify, is_system, status, "
    "is_shared, report_config, updated_at"
)


def _format_updated(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    text = str(value).replace("T", " ")
    if "." in text:
        text = text.split(".", 1)[0]
    return text[:19]


def _public_row(row: dict[str, Any]) -> dict[str, Any]:
    config = row.get("report_config")
    if isinstance(config, str):
        config = json.loads(config) if config else {}
    return {
        "id": row.get("id"),
        "uuid": row.get("uuid"),
        "code": row.get("code"),
        "name": row.get("name"),
        "description": row.get("description"),
        "category": row.get("category"),
        "classify": row.get("classify"),
        "is_system": bool(row.get("is_system")),
        "status": row.get("status"),
        "is_shared": bool(row.get("is_shared")),
        "updated_at": _format_updated(row.get("updated_at")),
        "report_config": config if isinstance(config, dict) else {},
    }


class TortoiseReportCenterStore:
    """读写已有表 apps_kuaireport_reports。不读数据集，不执行报表 SQL。"""

    def __init__(self, conn: Any) -> None:
        self.conn = conn

    async def _dicts(self, sql: str, params: list[Any]) -> list[dict[str, Any]]:
        if not hasattr(self.conn, "execute_query_dict"):
            raise RuntimeError("报表中心需要连接的 execute_query_dict")
        rows = await self.conn.execute_query_dict(sql, params)
        return [dict(row) for row in rows]

    async def list_rows(
        self,
        tenant_id: int,
        status: str | None,
        category: str | None,
        classify: str | None,
    ) -> list[dict[str, Any]]:
        sql = (
            f"SELECT {_REPORT_LIST_COLUMNS} "
            "FROM apps_kuaireport_reports WHERE tenant_id = $1"
        )
        params: list[Any] = [tenant_id]
        if status is not None:
            params.append(status)
            sql += f" AND status = ${len(params)}"
        if category is not None:
            params.append(category)
            sql += f" AND category = ${len(params)}"
        if classify is not None:
            params.append(classify)
            sql += f" AND classify = ${len(params)}"
        sql += " ORDER BY classify, code"
        return [_public_row(row) for row in await self._dicts(sql, params)]

    async def get_row(self, tenant_id: int, report_id: int) -> dict[str, Any] | None:
        rows = await self._dicts(
            f"SELECT {_REPORT_LIST_COLUMNS} "
            "FROM apps_kuaireport_reports WHERE tenant_id = $1 AND id = $2",
            [tenant_id, report_id],
        )
        if not rows:
            return None
        return _public_row(rows[0])

    async def existing_codes(self, tenant_id: int, codes: list[str]) -> set[str]:
        if not codes:
            return set()
        placeholders = ", ".join(f"${index + 2}" for index in range(len(codes)))
        rows = await self._dicts(
            "SELECT code FROM apps_kuaireport_reports "
            f"WHERE tenant_id = $1 AND code IN ({placeholders})",
            [tenant_id, *codes],
        )
        return {str(row["code"]) for row in rows}

    async def set_status(self, tenant_id: int, report_id: int, status: str) -> None:
        await self.conn.execute_query(
            "UPDATE apps_kuaireport_reports SET status = $3, updated_at = CURRENT_TIMESTAMP "
            "WHERE tenant_id = $1 AND id = $2",
            [tenant_id, report_id, status],
        )

    async def withdraw(self, tenant_id: int, report_id: int) -> dict[str, Any] | None:
        """回到草稿，并清掉分享令牌，已发出的链接立即失效。"""
        row = await self.get_row(tenant_id, report_id)
        if row is None:
            return None
        await self.conn.execute_query(
            "UPDATE apps_kuaireport_reports SET status = $3, is_shared = FALSE, "
            "share_token = NULL, share_expires_at = NULL, share_password_hash = NULL, "
            "share_allow_ip_cidrs = NULL, updated_at = CURRENT_TIMESTAMP "
            "WHERE tenant_id = $1 AND id = $2",
            [tenant_id, report_id, STATUS_DRAFT],
        )
        return await self.get_row(tenant_id, report_id)

    async def delete_row(self, tenant_id: int, report_id: int) -> bool:
        row = await self.get_row(tenant_id, report_id)
        if row is None:
            return False
        await self.conn.execute_query(
            "DELETE FROM apps_kuaireport_reports WHERE tenant_id = $1 AND id = $2",
            [tenant_id, report_id],
        )
        return True

    async def registered_dataset_uuid(self, tenant_id: int) -> str | None:
        """本租户第一条可用 dataset 数据源；种入报表只绑数据集，不绑 static/http。"""
        rows = await self._dicts(
            "SELECT uuid FROM apps_kuaireport_data_sources "
            "WHERE tenant_id = $1 AND type = 'dataset' "
            "ORDER BY is_default DESC, id LIMIT 1",
            [tenant_id],
        )
        if not rows:
            return None
        value = rows[0].get("uuid")
        if value is None or str(value).strip() == "":
            return None
        return str(value)

    async def bind_data_source(
        self, tenant_id: int, report_id: int, source_uuid: str
    ) -> dict[str, Any] | None:
        """把 report_config.extra.data_source_uuid 补写成 source_uuid，返回新行。"""
        row = await self.get_row(tenant_id, report_id)
        if row is None:
            return None
        config = dict(row.get("report_config") or {})
        extra = dict(config.get("extra") or {})
        extra[REPORT_DATA_SOURCE_UUID] = source_uuid
        config["extra"] = extra
        await self.conn.execute_query(
            "UPDATE apps_kuaireport_reports "
            "SET report_config = $3::jsonb, updated_at = CURRENT_TIMESTAMP "
            "WHERE tenant_id = $1 AND id = $2",
            [tenant_id, report_id, json.dumps(config, ensure_ascii=False)],
        )
        return await self.get_row(tenant_id, report_id)

    async def insert_system_report(self, row: dict[str, Any]) -> None:
        await self.conn.execute_query(
            "INSERT INTO apps_kuaireport_reports ("
            "uuid, tenant_id, code, name, category, classify, is_system, status, "
            "is_shared, report_config"
            ") VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10::jsonb)",
            [
                row["uuid"],
                row["tenant_id"],
                row["code"],
                row["name"],
                row["category"],
                row["classify"],
                row["is_system"],
                row["status"],
                row["is_shared"],
                json.dumps(row["report_config"], ensure_ascii=False),
            ],
        )


@asynccontextmanager
async def report_center_transaction() -> AsyncIterator[Any]:
    from tortoise.transactions import in_transaction

    async with in_transaction() as conn:
        yield TortoiseReportCenterStore(conn)


def _check_category(category: str | None) -> None:
    if category is None:
        return
    if category not in ALLOWED_CATEGORIES:
        raise ValidationError("category 只能是 system 或 custom")


def _bound_uuid(row: dict[str, Any]) -> str | None:
    config = row.get("report_config")
    if not isinstance(config, dict):
        return None
    extra = config.get("extra")
    if not isinstance(extra, dict):
        return None
    raw = extra.get(REPORT_DATA_SOURCE_UUID)
    if not isinstance(raw, str) or not raw.strip():
        return None
    return raw.strip()


async def _rebind_missing_sources(
    store: Any, tenant_id: int, rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """系统报表缺 data_source_uuid 时幂等回绑本租户第一条 dataset 数据源。"""
    missing_ids = {
        row["id"]
        for row in rows
        if row.get("is_system") and row.get("id") is not None and _bound_uuid(row) is None
    }
    if not missing_ids:
        return rows
    source_uuid = await store.registered_dataset_uuid(tenant_id)
    if not source_uuid:
        return rows
    rebound: dict[int, dict[str, Any]] = {}
    for report_id in missing_ids:
        saved = await store.bind_data_source(tenant_id, report_id, source_uuid)
        if saved is not None:
            rebound[report_id] = saved
    return [rebound.get(row["id"], row) for row in rows]


async def _keep_visible(
    store: Any,
    tenant_id: int,
    user_id: int | None,
    resource_type: str,
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    from apps.kuaireport.slices.s149_distribution import filter_visible_rows

    return await filter_visible_rows(tenant_id, user_id, resource_type, rows, store)


async def _assert_visible(
    store: Any, tenant_id: int, user_id: int | None, resource_type: str, resource_id: int
) -> None:
    from apps.kuaireport.slices.s149_distribution import enforce_grant_view, grant_reader

    reader = grant_reader(store)
    if reader is None:
        return
    await enforce_grant_view(tenant_id, user_id, resource_type, resource_id, store=reader)


async def list_reports(
    *,
    status: str | None = None,
    category: str | None = None,
    classify: str | None = None,
    user_id: int | None = None,
) -> list[dict[str, Any]]:
    tenant_id = require_tenant_id()
    _check_category(category)
    async with report_center_transaction() as store:
        rows = await store.list_rows(tenant_id, status, category, classify)
        rows = await _rebind_missing_sources(store, tenant_id, rows)
        return await _keep_visible(store, tenant_id, user_id, "report", rows)


async def get_report(report_id: int, *, user_id: int | None = None) -> dict[str, Any]:
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        row = await store.get_row(tenant_id, report_id)
        if row is None:
            raise NotFoundError("报表", str(report_id))
        await _assert_visible(store, tenant_id, user_id, "report", report_id)
    return row


def _seed_row(
    tenant_id: int,
    code: str,
    name: str,
    classify: str | None,
    data_source_uuid: str | None,
) -> dict[str, Any]:
    if name in FORBIDDEN_REPORT_NAMES:
        raise ValidationError("不能把菜单分组种成报表")
    return {
        "uuid": str(uuid.uuid4()),
        "tenant_id": tenant_id,
        "code": code,
        "name": name,
        "category": CATEGORY_SYSTEM,
        "classify": resolve_classify(classify),
        "is_system": True,
        "status": STATUS_DRAFT,
        "is_shared": False,
        "report_config": system_report_config(data_source_uuid),
    }


async def seed_system_reports(data_source_uuid: str | None = None) -> list[str]:
    """同一租户的十张在一个事务里种入。已有 (tenant_id, code) 不重复插入。"""
    tenant_id = require_tenant_id()
    codes = [code for code, _name, _classify in SYSTEM_REPORTS]
    inserted: list[str] = []
    async with report_center_transaction() as store:
        existing = await store.existing_codes(tenant_id, codes)
        for code, name, classify in SYSTEM_REPORTS:
            if code in existing:
                continue
            await store.insert_system_report(
                _seed_row(tenant_id, code, name, classify, data_source_uuid)
            )
            inserted.append(code)
    return inserted


async def _registered_data_source_uuid(tenant_id: int) -> str | None:
    from tortoise import Tortoise

    conn = Tortoise.get_connection("default")
    rows = await conn.execute_query_dict(
        "SELECT uuid FROM apps_kuaireport_data_sources "
        "WHERE tenant_id = $1 AND type = 'dataset' "
        "ORDER BY is_default DESC, id LIMIT 1",
        [tenant_id],
    )
    if not rows:
        return None
    value = rows[0].get("uuid")
    if value is None or str(value).strip() == "":
        return None
    return str(value)


async def seed_system_reports_on_lifecycle(tenant_id: int) -> list[str]:
    """安装或启用入口调用。无租户上下文不种入。"""
    current = get_current_tenant_id()
    if current is None or int(current) != int(tenant_id):
        return []
    source_uuid = await _registered_data_source_uuid(int(current))
    return await seed_system_reports(source_uuid)


async def publish_report(report_id: int) -> dict[str, Any]:
    """把 status 写成 PUBLISHED。"""
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        row = await store.get_row(tenant_id, report_id)
        if row is None:
            raise NotFoundError("报表", str(report_id))
        await store.set_status(tenant_id, report_id, STATUS_PUBLISHED)
        saved = await store.get_row(tenant_id, report_id)
    if saved is None:
        raise NotFoundError("报表", str(report_id))
    return saved


def report_menu_path(report_id: int) -> str:
    return f"/apps/kuaireport/reports/view/{report_id}"


async def withdraw_report(report_id: int) -> dict[str, Any]:
    """撤回为草稿，并让已有分享链接失效。"""
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        saved = await store.withdraw(tenant_id, report_id)
    if saved is None:
        raise NotFoundError("报表", str(report_id))
    return saved


async def delete_report(report_id: int) -> None:
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        deleted = await store.delete_row(tenant_id, report_id)
    if not deleted:
        raise NotFoundError("报表", str(report_id))
    await clear_report_mount(report_id)


class MountReportBody(BaseModel):
    parent_uuid: str = Field(..., min_length=1)
    menu_name: str = Field(..., min_length=1, max_length=100)


async def _report_menu(tenant_id: int, report_id: int):
    from core.models.menu import Menu

    return await Menu.filter(
        tenant_id=tenant_id,
        path=report_menu_path(report_id),
        deleted_at__isnull=True,
    ).first()


async def _parent_uuid_of(menu) -> str | None:
    from core.models.menu import Menu

    if menu is None or not menu.parent_id:
        return None
    parent = await Menu.get_or_none(id=menu.parent_id, deleted_at__isnull=True)
    if parent is None:
        return None
    return str(parent.uuid)


def _mount_payload(menu, parent_uuid: str | None) -> dict[str, Any]:
    if menu is None:
        return {"mounted": False, "menu_uuid": None, "menu_name": None, "parent_uuid": None}
    return {
        "mounted": True,
        "menu_uuid": str(menu.uuid),
        "menu_name": menu.name,
        "parent_uuid": parent_uuid,
    }


async def get_report_mount(report_id: int) -> dict[str, Any]:
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        row = await store.get_row(tenant_id, report_id)
    if row is None:
        raise NotFoundError("报表", str(report_id))
    menu = await _report_menu(tenant_id, report_id)
    return _mount_payload(menu, await _parent_uuid_of(menu))


async def mount_report(report_id: int, parent_uuid: str, menu_name: str) -> dict[str, Any]:
    """把报表挂到应用菜单下。系统菜单（没有 application_uuid）不可作为父级。"""
    from core.models.menu import Menu
    from core.schemas.menu import MenuCreate, MenuUpdate
    from core.services.system.menu_service import MenuService

    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        row = await store.get_row(tenant_id, report_id)
    if row is None:
        raise NotFoundError("报表", str(report_id))
    parent = await Menu.filter(
        uuid=parent_uuid,
        tenant_id=tenant_id,
        deleted_at__isnull=True,
    ).first()
    if parent is None:
        raise ValidationError("父菜单不存在或不属于当前组织")
    if not (parent.application_uuid or "").strip():
        raise ValidationError("系统菜单不可选")
    name = menu_name.strip()
    if not name:
        raise ValidationError("请填写菜单名称")
    path = report_menu_path(report_id)
    existing = await _report_menu(tenant_id, report_id)
    if existing is None:
        await MenuService.create_menu(
            tenant_id,
            MenuCreate(
                name=name,
                path=path,
                icon="barChart",
                parent_uuid=str(parent.uuid),
                permission_code="kuaireport:report:display",
                sort_order=100,
            ),
        )
    else:
        await MenuService.update_menu(
            tenant_id,
            str(existing.uuid),
            MenuUpdate(name=name, parent_uuid=str(parent.uuid)),
        )
    menu = await _report_menu(tenant_id, report_id)
    return _mount_payload(menu, await _parent_uuid_of(menu))


async def clear_report_mount(report_id: int) -> dict[str, Any]:
    from core.services.system.menu_service import MenuService

    tenant_id = require_tenant_id()
    menu = await _report_menu(tenant_id, report_id)
    if menu is not None:
        await MenuService.delete_menu(tenant_id, str(menu.uuid))
    return {"mounted": False, "menu_uuid": None, "menu_name": None, "parent_uuid": None}


def _is_secret_key(key: str) -> bool:
    return _SECRET_KEY.search(key) is not None


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str):
        if _SECRET_VALUE.search(value) or _PATH_VALUE.search(value):
            return None
        if value.startswith(_FORMULA_PREFIXES):
            return "'" + value
        return value
    if isinstance(value, (dict, list)):
        return None
    return value


def _export_columns(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key in seen or _is_secret_key(str(key)):
                continue
            seen.add(key)
            columns.append(str(key))
    return columns


def build_full_excel(name: str, rows: list[dict[str, Any]], summary: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (name or "报表")[:31]
    columns = _export_columns(rows)
    sheet.append(columns)
    for row in rows:
        sheet.append([_safe_cell(row.get(column)) for column in columns])
    if summary:
        sheet.append([])
        sheet.append(["合计"])
        for key, value in summary.items():
            if _is_secret_key(str(key)):
                continue
            sheet.append([str(key), _safe_cell(value)])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _execute_page(result: Any) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
    """读 ExecuteReportResult 的字段。属性或 model_dump()，不用 dict.get。"""
    if hasattr(result, "model_dump"):
        payload = result.model_dump()
        data = payload["data"]
        total = payload["total"]
        summary = payload["summary"]
    else:
        data = result.data
        total = result.total
        summary = result.summary
    page = list(data or [])
    summary_rows = summary if isinstance(summary, dict) else {}
    return page, int(total or 0), summary_rows


def _safe_filename(code: Any) -> str:
    """Content-Disposition 文件名：去掉引号/CR-LF/非 ASCII，兜底 'report'。"""
    text = _SAFE_FILENAME_CHARS.sub("_", str(code or ""))
    text = text.strip("._ ")
    if not text:
        text = "report"
    return f"{text}.xlsx"


async def export_full_excel(
    report_id: int,
    filters: dict[str, Any] | None,
    *,
    viewer_id: int | None = None,
    apply_grant: bool = False,
) -> tuple[bytes, str]:
    """按同一套筛选翻页调用 execute_report，写成 xlsx。忽略调用方自带的 limit/offset。"""
    tenant_id = require_tenant_id()
    async with report_center_transaction() as store:
        report = await store.get_row(tenant_id, report_id)
        if report is None:
            raise NotFoundError("报表", str(report_id))
        if apply_grant:
            await _assert_visible(store, tenant_id, viewer_id, "report", report_id)

    execute_report = resolve_execute_report()
    base = {
        key: value
        for key, value in (filters or {}).items()
        if key not in ("limit", "offset")
    }
    rows: list[dict[str, Any]] = []
    summary: dict[str, Any] = {}
    offset = 0
    while True:
        page_filters = {**base, "limit": FULL_EXPORT_PAGE_SIZE, "offset": offset}
        result = await execute_report(tenant_id, report_id, page_filters)
        page, _total, page_summary = _execute_page(result)
        if page_summary:
            summary = page_summary
        rows.extend(page)
        offset += len(page)
        if len(page) < FULL_EXPORT_PAGE_SIZE:
            break
    filename = _safe_filename(report["code"])
    return build_full_excel(str(report["name"]), rows, summary), filename


@router.get(
    "",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def list_reports_api(
    status: str | None = None,
    category: str | None = None,
    classify: str | None = None,
    tenant_id: int = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
) -> list[dict[str, Any]]:
    return await list_reports(
        status=status, category=category, classify=classify, user_id=current_user.id
    )


@router.get(
    "/{report_id:int}",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def get_report_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    return await get_report(report_id, user_id=current_user.id)


@router.post(
    "/{report_id:int}/publish",
    dependencies=[Depends(require_permission_codes("kuaireport:report:publish"))],
)
async def publish_report_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, Any]:
    return await publish_report(report_id)


@router.post(
    "/{report_id:int}/withdraw",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def withdraw_report_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, Any]:
    return await withdraw_report(report_id)


@router.delete(
    "/{report_id:int}",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def delete_report_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, bool]:
    await delete_report(report_id)
    return {"success": True}


@router.get(
    "/{report_id:int}/mount",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def get_report_mount_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, Any]:
    return await get_report_mount(report_id)


@router.post(
    "/{report_id:int}/mount",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def mount_report_api(
    report_id: int,
    body: MountReportBody,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, Any]:
    return await mount_report(report_id, body.parent_uuid, body.menu_name)


@router.delete(
    "/{report_id:int}/mount",
    dependencies=[Depends(require_permission_codes("kuaireport:report:display"))],
)
async def clear_report_mount_api(
    report_id: int,
    tenant_id: int = Depends(get_current_tenant),
) -> dict[str, Any]:
    return await clear_report_mount(report_id)


@router.post(
    "/{report_id:int}/excel",
    dependencies=[Depends(require_permission_codes("kuaireport:report:export"))],
)
async def export_report_excel_api(
    report_id: int,
    body: FullExcelRequest,
    tenant_id: int = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
) -> Response:
    content, filename = await export_full_excel(
        report_id, body.filters, viewer_id=current_user.id, apply_grant=True
    )
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
