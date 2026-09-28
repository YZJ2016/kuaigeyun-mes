"""spec 146：报表中心列表、十张系统报表种入、后端全量 Excel。"""

from __future__ import annotations

import json
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.kuaireport.slices import s146_center as center
from infra.domain.tenant_context import TenantContextError, clear_tenant_context, set_current_tenant_id

TENANT = 9
FRONTEND_REPORTS = (
    Path(__file__).resolve().parents[4]
    / "riveredge-frontend"
    / "src"
    / "apps"
    / "kuaireport"
    / "pages"
    / "reports"
)


class MemoryStore:
    def __init__(self) -> None:
        self.rows: list[dict] = []
        self._pending: list[dict] | None = None
        self.fail_on_insert: int | None = None
        self.inserts = 0
        self.next_id = 1

    def _visible(self) -> list[dict]:
        if self._pending is None:
            return list(self.rows)
        return [*self.rows, *self._pending]

    async def list_rows(self, tenant_id, status, category, classify):
        rows = []
        for row in self._visible():
            if row["tenant_id"] != tenant_id:
                continue
            if status is not None and row["status"] != status:
                continue
            if category is not None and row["category"] != category:
                continue
            if classify is not None and row["classify"] != classify:
                continue
            rows.append(center._public_row(row))
        rows.sort(key=lambda item: (item["classify"], item["code"]))
        return rows

    async def get_row(self, tenant_id, report_id):
        for row in self._visible():
            if row["tenant_id"] == tenant_id and row["id"] == report_id:
                return center._public_row(row)
        return None

    async def existing_codes(self, tenant_id, codes):
        wanted = set(codes)
        return {
            row["code"]
            for row in self._visible()
            if row["tenant_id"] == tenant_id and row["code"] in wanted
        }

    async def set_status(self, tenant_id, report_id, status):
        for bucket in (self._pending or [], self.rows):
            for row in bucket:
                if row["tenant_id"] == tenant_id and row["id"] == report_id:
                    row["status"] = status
                    return

    async def insert_system_report(self, row):
        self.inserts += 1
        if self.fail_on_insert is not None and self.inserts == self.fail_on_insert:
            raise RuntimeError("seed failed")
        stored = dict(row)
        stored["id"] = self.next_id
        self.next_id += 1
        assert self._pending is not None
        self._pending.append(stored)


@asynccontextmanager
async def _open(store: MemoryStore):
    store._pending = []
    try:
        yield store
    except Exception:
        store._pending = None
        raise
    else:
        store.rows.extend(store._pending)
        store._pending = None


@pytest.fixture
def store(monkeypatch):
    memory = MemoryStore()

    @asynccontextmanager
    async def report_center_transaction():
        async with _open(memory) as current:
            yield current

    monkeypatch.setattr(center, "report_center_transaction", report_center_transaction)
    set_current_tenant_id(TENANT)
    try:
        yield memory
    finally:
        clear_tenant_context()


def test_system_report_codes_match_spec():
    codes = [code for code, _name, _classify in center.SYSTEM_REPORTS]
    assert codes == [
        "inv_ledger",
        "wo_tracking",
        "qc_pass_rate",
        "perf_stats",
        "node_timeliness",
        "process_efficiency",
        "sales_trace",
        "purchase_trace",
        "material_trace",
        "biz_overview",
    ]
    names = {code: name for code, name, _classify in center.SYSTEM_REPORTS}
    assert names["sales_trace"] == "销售全链路追踪"
    assert names["purchase_trace"] == "采购全链路追踪"
    assert center.FORBIDDEN_REPORT_NAMES.isdisjoint(names.values())


@pytest.mark.asyncio
async def test_no_tenant_does_not_read_reports():
    clear_tenant_context()
    with pytest.raises(TenantContextError):
        await center.list_reports()


@pytest.mark.asyncio
async def test_seed_is_idempotent_and_defaults(store: MemoryStore):
    inserted = await center.seed_system_reports()
    assert inserted == [code for code, _name, _classify in center.SYSTEM_REPORTS]
    assert await center.seed_system_reports() == []
    rows = await center.list_reports()
    assert len(rows) == 10
    by_code = {row["code"]: row for row in rows}
    assert by_code["sales_trace"]["classify"] == "销售"
    assert by_code["purchase_trace"]["classify"] == "采购"
    for code, _name, classify in center.SYSTEM_REPORTS:
        row = by_code[code]
        assert row["category"] == "system"
        assert row["is_system"] is True
        assert row["status"] == "DRAFT"
        assert row["is_shared"] is False
        assert row["name"] not in center.FORBIDDEN_REPORT_NAMES
        assert row["classify"] == classify
        blob = json.dumps(row["report_config"], ensure_ascii=False).lower()
        assert "select " not in blob
        assert "query_config" not in blob
        assert "data_source_id" not in row["report_config"]
        assert row["report_config"]["extra"]["data_source_uuid"] is None
        assert "dataset_uuid" not in row["report_config"]
        assert row["report_config"]["fields"] == []
    assert await center.list_reports(category="custom") == []


@pytest.mark.asyncio
async def test_seed_rolls_back_when_insert_fails(store: MemoryStore):
    store.fail_on_insert = 4
    with pytest.raises(RuntimeError, match="seed failed"):
        await center.seed_system_reports()
    assert store.rows == []
    store.fail_on_insert = None
    store.inserts = 0
    inserted = await center.seed_system_reports()
    assert len(inserted) == 10


@pytest.mark.asyncio
async def test_list_filters_status_category_and_classify(store: MemoryStore):
    await center.seed_system_reports()
    store.rows.append(
        {
            "id": 100,
            "uuid": "custom-1",
            "tenant_id": TENANT,
            "code": "my_sheet",
            "name": "我的报表",
            "category": "custom",
            "classify": "销售",
            "is_system": False,
            "status": "NOT_DRAFT_SAMPLE",
            "is_shared": False,
            "report_config": {"dataset_code": "my_sheet", "fields": []},
        }
    )
    drafts = await center.list_reports(status="DRAFT")
    assert len(drafts) == 10
    assert all(row["status"] == "DRAFT" for row in drafts)
    other = await center.list_reports(status="NOT_DRAFT_SAMPLE")
    assert [row["code"] for row in other] == ["my_sheet"]
    custom = await center.list_reports(category="custom")
    assert [row["code"] for row in custom] == ["my_sheet"]
    sales = await center.list_reports(classify="销售")
    assert {row["code"] for row in sales} == {"sales_trace", "my_sheet"}
    with pytest.raises(center.ValidationError):
        await center.list_reports(category="analysis")


@pytest.mark.asyncio
async def test_full_excel_pages_through_execute_report(store: MemoryStore, monkeypatch):
    await center.seed_system_reports()
    report = next(row for row in store.rows if row["code"] == "inv_ledger")
    calls: list[dict] = []

    class _ExecutePage:
        def __init__(self, data, total, summary):
            self.data = data
            self.total = total
            self.summary = summary

        def model_dump(self):
            return {"data": self.data, "total": self.total, "summary": self.summary}

    async def fake_execute(tenant_id, report_id, filters):
        assert tenant_id == TENANT
        assert report_id == report["id"]
        calls.append(dict(filters))
        all_rows = [
            {"qty": index, "note": f"n{index}", "password": "hidden", "file": r"C:\secret\a.xlsx"}
            for index in range(5)
        ]
        offset = filters["offset"]
        limit = filters["limit"]
        return _ExecutePage(
            all_rows[offset : offset + limit],
            5,
            {"qty": 10, "token": "hidden-token"},
        )

    monkeypatch.setattr(center, "FULL_EXPORT_PAGE_SIZE", 2)
    monkeypatch.setattr(center, "resolve_execute_report", lambda: fake_execute)
    content, filename = await center.export_full_excel(
        report["id"],
        {"warehouse": "A", "limit": 1, "offset": 99},
    )
    assert filename == "inv_ledger.xlsx"
    assert content[:2] == b"PK"
    assert [item["offset"] for item in calls] == [0, 2, 4]
    assert all(item["limit"] == 2 and item["warehouse"] == "A" for item in calls)
    assert all("offset" in item and item["offset"] != 99 for item in calls)
    sheet = load_workbook(BytesIO(content)).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "qty" in headers and "note" in headers
    assert "password" not in headers
    qty_at = headers.index("qty")
    note_at = headers.index("note")
    data_rows = list(sheet.iter_rows(min_row=2, max_row=6))
    assert [row[qty_at].value for row in data_rows] == list(range(5))
    assert [row[note_at].value for row in data_rows] == [f"n{index}" for index in range(5)]
    flat = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "hidden" not in flat
    assert "hidden-token" not in flat
    assert r"C:\secret\a.xlsx" not in flat
    assert "合计" in flat
    assert "qty" in flat


@pytest.mark.asyncio
async def test_publish_writes_published_status(store: MemoryStore):
    await center.seed_system_reports("source-uuid-1")
    report = next(row for row in store.rows if row["code"] == "inv_ledger")
    assert report["report_config"]["extra"]["data_source_uuid"] == "source-uuid-1"
    published = await center.publish_report(report["id"])
    assert published["status"] == "PUBLISHED"
    assert published["status"] == center.STATUS_PUBLISHED


@pytest.mark.asyncio
async def test_lifecycle_seed_skips_without_tenant_context():
    clear_tenant_context()
    assert await center.seed_system_reports_on_lifecycle(TENANT) == []


def test_slice_does_not_implement_a_second_executor():
    source = Path(center.__file__).read_text(encoding="utf-8")
    assert "def execute_report" not in source
    assert "core_datasets" not in source
    assert "useUniReportExport" not in source
    assert "exportDomainReport" not in source


def test_preview_page_uses_unireport_not_browser_export():
    preview = (FRONTEND_REPORTS / "preview.tsx").read_text(encoding="utf-8")
    assert "UniReport" in preview
    assert 'mode="config"' in preview
    assert "showPrintButton" in preview
    assert "showExportButton={false}" in preview
    assert "useUniReportExport" not in preview
    assert "exportDomainReport" not in preview
    joined = "\n".join(path.read_text(encoding="utf-8") for path in FRONTEND_REPORTS.glob("*.tsx"))
    assert "效能中心" not in joined
    assert "ReportSharedView" not in joined
    assert "DashboardSharedView" not in joined
